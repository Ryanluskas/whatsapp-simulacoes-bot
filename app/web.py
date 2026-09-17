"""API HTTP do painel, com tempo real por SSE.

O tempo real deixou de usar Socket.IO e passou a usar *Server-Sent Events*.
Motivos concretos: o fluxo aqui e' so' servidor -> navegador, o ``EventSource``
e' nativo (nenhum javascript de terceiro para carregar de um CDN, o que fazia o
painel quebrar sem internet), ele reconecta sozinho e carrega o cookie de
sessao como qualquer requisicao HTTP - o que resolve a autenticacao de graca.

Correcoes de seguranca em relacao a versao anterior:

* O canal em tempo real nao tinha autenticacao nenhuma: qualquer um que
  alcancasse a porta recebia o fluxo ao vivo com CPF. Agora ele e' uma rota
  protegida como todas as outras.
* ``MASK_CPF_IN_UI`` existia no ``.env`` e nunca era aplicado - o CPF completo
  saia na API e nos exports. Agora o mascaramento acontece no servidor.
* ``SESSION_SECRET`` tambem nunca era usado; as sessoes viviam num ``set`` em
  memoria e caiam a cada reinicio. Agora sao tokens assinados por HMAC.
* Login sem limite de tentativas.
* Os arquivos temporarios de XLSX/PDF ficavam para tras a cada export.
"""

from __future__ import annotations

import asyncio
import hmac
import csv
import io
import json
import os
import re
import tempfile
from contextlib import asynccontextmanager, suppress

from pathlib import Path
from typing import Any

from fastapi import BackgroundTasks, Depends, FastAPI, Form, HTTPException, Query, Request
from fastapi.responses import FileResponse, JSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles

from . import analytics
from .clock import now_iso, range_bounds
from .evolution_webhook import interpretar_todos
from .config import Config
from .db import Database
from .events import EventHub
from .manager import COMPROVANTES_DIR
from .models import STAGE_LABELS, STATUS_LABELS
from .security import (
    SESSION_COOKIE,
    RateLimiter,
    SessionManager,
    check_password,
)

STATIC_DIR = Path(__file__).resolve().parent.parent / "dashboard" / "static"

EXPORT_COLUMNS = [
    ("request_id", "Solicitação"),
    ("created_at", "Criada em"),
    ("consultant_name", "Consultor"),
    ("customer_name", "Cliente"),
    ("cpf_display", "CPF"),
    ("bank", "Banco"),
    ("contract", "Contrato"),
    ("status_label", "Status"),
    ("refin", "Refinanciamento"),
    ("reduction_value", "Valor disponível"),
    ("installment_sum", "Soma parcelas"),
    ("installment_count", "Qtd parcelas"),
    ("debt_sum", "Saldo devedor"),
    ("margin", "Margem"),
    ("attempts", "Tentativas"),
    ("processing_seconds", "Tempo (s)"),
    ("finished_at", "Concluída em"),
    ("error_message", "Erro"),
]


def create_app(config: Config, db: Database, hub: EventHub, manager) -> FastAPI:
    broadcaster = SseBroadcaster(hub)

    @asynccontextmanager
    async def lifespan(app: FastAPI):
        # Uma unica ponte entre as threads (WhatsApp, fila, simulador) e o
        # mundo asyncio. Cada aba conectada e' apenas mais um destino do fan-out.
        await broadcaster.start()
        try:
            yield
        finally:
            await broadcaster.stop()

    api = FastAPI(
        title="Allana · Central de Simulações",
        docs_url=None,
        redoc_url=None,
        lifespan=lifespan,
    )
    sessions = SessionManager(config.session_secret, config.session_hours)
    login_limiter = RateLimiter(max_attempts=8, window_seconds=300)
    mask = config.mask_cpf_in_ui

    # --------------------------------------------------------------- auth
    def current_user(request: Request) -> str | None:
        return sessions.verify(request.cookies.get(SESSION_COOKIE))

    def require_auth(request: Request) -> str:
        user = current_user(request)
        if not user:
            raise HTTPException(status_code=401, detail="não autenticado")
        return user

    Auth = Depends(require_auth)

    @api.post("/api/login")
    async def login(request: Request, password: str = Form(...)):
        client = request.client.host if request.client else "desconhecido"
        if not login_limiter.allow(client):
            raise HTTPException(
                status_code=429,
                detail=f"muitas tentativas, aguarde {login_limiter.retry_after(client)}s",
            )
        if not check_password(password, config.dashboard_password):
            manager.log("WARNING", "auth", f"Tentativa de login inválida de {client}.")
            raise HTTPException(status_code=401, detail="senha inválida")
        login_limiter.reset(client)
        token = sessions.issue("admin")
        manager.log("INFO", "auth", f"Login realizado de {client}.")
        response = JSONResponse({"ok": True, "user": "admin"})
        response.set_cookie(
            SESSION_COOKIE, token, httponly=True, samesite="lax",
            max_age=sessions.ttl_seconds, path="/",
        )
        return response

    @api.post("/api/logout")
    async def logout(request: Request):
        sessions.revoke(request.cookies.get(SESSION_COOKIE))
        response = JSONResponse({"ok": True})
        response.delete_cookie(SESSION_COOKIE, path="/")
        return response

    @api.get("/api/health")
    async def health():
        """Sinal de vida para o iniciar.bat e para qualquer supervisor.

        Sem autenticação de propósito - é o que o script consulta antes de
        dizer que o sistema subiu. Por isso devolve só estado de serviço,
        nunca dado de cliente.
        """
        try:
            whatsapp = manager.whatsapp.status.state
        except Exception:
            whatsapp = "unknown"
        simulators = [s for s in manager.simulators if getattr(s, "running", False)]
        servicos = {
            "web": True,
            "whatsapp": whatsapp,
            "simulators_running": len(simulators),
            "simulators_total": len(manager.simulators),
            "queue_depth": manager.queue.depth(),
        }
        # Diagnostico da Evolution: so' booleanos e estado, para responder
        # "por que está ligado e não responde?" sem sessão e sem expor nada.
        try:
            diagnostico = manager.diagnostico_do_whatsapp().get("evolution")
        except Exception:
            diagnostico = None
        if diagnostico:
            servicos["evolution"] = {
                chave: diagnostico.get(chave) for chave in (
                    "evolution_api_reachable", "api_key_valid", "instance_found",
                    "evolution_state", "group_configured", "webhook_configured",
                    "webhook_points_to_bot", "webhook_token_configured",
                    "last_webhook_at", "checked_at")
            }
        return {"status": "ok", "services": servicos}

    @api.get("/api/session")
    async def session(request: Request):
        user = current_user(request)
        return {
            "authenticated": bool(user),
            "user": user,
            "mask_cpf": mask,
            "group_name": config.whatsapp_group_name,
        }

    # ------------------------------------------------------------ dashboard
    @api.get("/api/bootstrap")
    async def bootstrap(_: str = Auth):
        return {
            "metrics": manager.metrics_snapshot(),
            "whatsapp": manager.whatsapp_status_payload(),
            "queue": manager.queue_snapshot(),
            "system": manager.system_status_payload(),
            "monitor": hub.recent(limit=150),
            "labels": {"stages": STAGE_LABELS, "statuses": STATUS_LABELS},
        }

    @api.get("/api/metrics")
    async def metrics(_: str = Auth):
        return manager.metrics_snapshot()

    @api.get("/api/system")
    async def system(_: str = Auth):
        return manager.system_status_payload()

    @api.get("/api/monitor")
    async def monitor(limit: int = Query(150, ge=1, le=500), _: str = Auth):
        return {"items": hub.recent(limit=limit)}

    @api.get("/api/queue")
    async def queue(_: str = Auth):
        return manager.queue_snapshot()

    # ------------------------------------------------------------- whatsapp
    @api.get("/api/whatsapp/status")
    async def whatsapp_status(_: str = Auth):
        return manager.whatsapp_status_payload()

    @api.post("/api/whatsapp/reconnect")
    async def whatsapp_reconnect(_: str = Auth):
        manager.reconnect_whatsapp()
        return {"ok": True}

    @api.get("/api/whatsapp/qr")
    async def whatsapp_qr(_: str = Auth):
        try:
            return {"qr": manager.qr_data_url(), "time": now_iso()}
        except Exception as exc:
            return {"qr": "", "error": str(exc)[:200]}

    # --------------------------------------------------------- simulacoes
    @api.get("/api/simulations")
    async def simulations(
        request: Request,
        limit: int = Query(50, ge=1, le=500),
        offset: int = Query(0, ge=0),
        _: str = Auth,
    ):
        filters = {k: v for k, v in request.query_params.items() if k not in {"limit", "offset"}}
        rows, total = analytics.list_simulations(db, filters, config.tz, limit=limit, offset=offset)
        return {
            "items": [analytics.serialize_simulation(r, mask) for r in rows],
            "total": total,
            "limit": limit,
            "offset": offset,
        }

    @api.get("/api/simulations/{sim_id}")
    async def simulation_detail(sim_id: int, _: str = Auth):
        row = db.fetchone("SELECT * FROM simulations WHERE id=?", (sim_id,))
        if not row:
            raise HTTPException(status_code=404, detail="simulação não encontrada")
        # A evidencia de cada envio vai junto: qual mensagem foi citada, qual
        # id o WhatsApp deu a resposta, a tentativa e o status HTTP. E' o que
        # permite reconstruir um incidente sem abrir o banco na mao.
        messages = db.fetchall(
            "SELECT id, direction, kind, text, media_path, status, created_at, consultant_name, "
            "       wa_message_id, provider, attempt, origin_message_id, quoted_message_id, "
            "       quote_status, quote_error, desfecho, http_status, media_id, error "
            "FROM messages WHERE simulation_id=? ORDER BY id ASC",
            (sim_id,),
        )
        # O caminho em disco nunca sai para o navegador; vira uma URL servida
        # pela rota autenticada de comprovantes.
        for mensagem in messages:
            caminho = mensagem.pop("media_path", None)
            mensagem["media_url"] = (
                f"/api/comprovantes/{Path(caminho).name}" if caminho else None
            )
        contracts: list[dict] = []
        if row.get("contracts_json"):
            try:
                contracts = json.loads(row["contracts_json"])
            except (ValueError, TypeError):
                contracts = []
        return {
            "simulation": analytics.serialize_simulation(row, mask),
            "timeline": analytics.timeline_for(db, row.get("request_id") or ""),
            "messages": messages,
            "contracts": contracts,
        }

    @api.post("/api/simulations/{sim_id}/entrega")
    async def resolver_entrega(sim_id: int, request: Request, usuario: str = Auth):
        """Entrega incerta: o operador olhou o grupo e diz se chegou.

        ``{"acao": "chegou"}`` fecha como entregue sem enviar nada;
        ``{"acao": "nao_chegou"}`` libera um reenvio. 409 quando a entrega
        ja' nao esta' incerta (outro clique ou outra aba chegou antes).
        """
        try:
            dados = await request.json()
        except ValueError:
            dados = {}
        acao = str((dados if isinstance(dados, dict) else {}).get("acao") or "")
        try:
            resultado = manager.resolver_entrega_incerta(sim_id, acao, quem=usuario)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        except LookupError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        if not resultado["ok"]:
            raise HTTPException(status_code=409, detail=resultado["motivo"])
        return resultado

    @api.get("/api/comprovantes/{nome}")
    async def comprovante(nome: str, _: str = Auth):
        """Serve a imagem enviada ao grupo.

        O nome e' validado contra um padrao fechado e o caminho final e'
        conferido contra a pasta: nao ha' como escapar dela com '..'.
        """
        if not re.fullmatch(r"REQ\d{1,12}\.png", nome):
            raise HTTPException(status_code=404, detail="não encontrado")
        pasta = Path(getattr(manager, "comprovantes_dir", COMPROVANTES_DIR))
        caminho = (pasta / nome).resolve()
        if not caminho.is_file() or caminho.parent != pasta.resolve():
            raise HTTPException(status_code=404, detail="não encontrado")
        return FileResponse(caminho, media_type="image/png")

    # ------------------------------------------------- agente do simulador
    # Usadas pelo agente que roda no Windows quando SIMULATOR_MODE=remote.
    # Autenticação por token próprio (AGENT_TOKEN), separado da sessão do
    # painel: são clientes diferentes, com poderes diferentes.
    def require_agent(request: Request) -> str:
        esperado = config.agent_token
        if not esperado:
            raise HTTPException(status_code=503, detail="AGENT_TOKEN não configurado")
        enviado = request.headers.get("x-agent-token", "")
        if not enviado or not hmac.compare_digest(enviado, esperado):
            raise HTTPException(status_code=401, detail="token do agente inválido")
        return request.headers.get("x-agent-name", "") or "agente"

    AgentAuth = Depends(require_agent)

    def _remotos():
        alvos = [s for s in manager.simulators if hasattr(s, "claim")]
        if not alvos:
            raise HTTPException(
                status_code=409,
                detail="o sistema não está em SIMULATOR_MODE=remote",
            )
        return alvos

    @api.post("/api/agent/claim")
    async def agent_claim(agente: str = AgentAuth):
        """Entrega a próxima simulação da fila, ou 204 se não houver."""
        for remoto in _remotos():
            tarefa = remoto.claim(agente)
            if tarefa:
                return tarefa
        return Response(status_code=204)

    @api.post("/api/agent/stage")
    async def agent_stage(request: Request, _: str = AgentAuth):
        dados = await request.json()
        request_id = str(dados.get("request_id") or "")
        stage = str(dados.get("stage") or "")
        aceito = any(r.report_stage(request_id, stage) for r in _remotos())
        return {"ok": aceito}

    @api.post("/api/agent/result")
    async def agent_result(request: Request, _: str = AgentAuth):
        dados = await request.json()
        request_id = str(dados.get("request_id") or "")
        if not request_id:
            raise HTTPException(status_code=400, detail="request_id ausente")
        aceito = any(r.submit_result(request_id, dados) for r in _remotos())
        if not aceito:
            # A simulação já expirou ou foi concluída por outra via.
            raise HTTPException(status_code=409, detail="solicitação não está mais aguardando")
        return {"ok": True}

    @api.get("/api/agent/ping")
    async def agent_ping(agente: str = AgentAuth):
        remotos = _remotos()   # levanta 409 se o modo remoto não estiver ativo
        return {
            "ok": True,
            "agent": agente,
            "time": now_iso(),
            "waiting": sum(r.agent_info()["waiting"] for r in remotos),
        }

    # ------------------------------------------------- webhook da Evolution
    # A entrada de mensagens quando WHATSAPP_MODE=evolution. No modo 'dom'
    # esta rota existe mas recusa: o bot esta lendo a tela, e aceitar mensagem
    # por dois caminhos ao mesmo tempo duplicaria toda solicitacao.
    #
    # Autenticacao propria, como a do agente. Ela recebe dado de cliente (nome
    # e CPF) e ENFILEIRA trabalho: deixa-la aberta seria dar a qualquer um na
    # rede o poder de mandar o bot simular o que quisesse.
    def require_webhook(request: Request) -> None:
        esperado = config.evolution_webhook_token
        if not esperado:
            raise HTTPException(status_code=503,
                                detail="EVOLUTION_WEBHOOK_TOKEN não configurado")
        enviado = (request.headers.get("x-webhook-token")
                   or request.headers.get("authorization", "").removeprefix("Bearer ").strip())
        if not enviado or not hmac.compare_digest(enviado, esperado):
            raise HTTPException(status_code=401, detail="token do webhook inválido")

    @api.post("/webhook/whatsapp")
    async def webhook_whatsapp(request: Request):
        require_webhook(request)

        if config.whatsapp_mode != "evolution":
            raise HTTPException(status_code=409,
                                detail="o sistema não está em WHATSAPP_MODE=evolution")
        try:
            payload = await request.json()
        except Exception:
            raise HTTPException(status_code=400, detail="corpo não é JSON")

        # A Evolution ALCANCOU o bot: e' o sinal que o diagnostico mostra.
        manager.registrar_webhook()
        leituras = interpretar_todos(payload, config.evolution_group_jid,
                                     config.whatsapp_group_name)

        if leituras[0].motivo == "connection.update":
            # A instancia caiu ou voltou. Refletir no painel na hora, em vez de
            # esperar o vigia: "conectado na tela e mudo no grupo" foi o pior
            # cenario da camada antiga.
            estado = str(((payload.get("data") or {}).get("state") or "")).lower()
            if estado and estado != "open":
                manager.log("ERROR", "whatsapp",
                            f"A instância da Evolution caiu (state={estado}).")
            return {"ok": True, "acao": "connection.update"}

        respostas = []
        for leitura in leituras:
            if not leitura:
                # 200 de proposito: ignorar nao e' erro, e devolver 4xx faria a
                # Evolution reentregar para sempre uma mensagem que nunca vamos
                # querer.
                respostas.append({"ok": True, "ignorado": leitura.motivo})
                continue

            if leitura.aviso:
                manager.log("WARNING", "whatsapp", leitura.aviso)

            # GRAVAR ANTES DE RESPONDER 200. A gravacao e' a trava contra
            # reentrega (no banco, sobrevive a reinicio) e e' o que permite
            # retomar a mensagem se o processo cair antes de trata-la. Se ela
            # falhar, a excecao vira 500 e a Evolution reentrega -- que e'
            # exatamente o comportamento certo.
            recebida = manager.receber_mensagem(leitura.mensagem)
            if recebida.get("aceita"):
                respostas.append({"ok": True, "request": "enfileirada",
                                  "message_id": leitura.mensagem.message_id})
            else:
                respostas.append({"ok": True, "ignorado": "já processado",
                                  "message_id": leitura.mensagem.message_id,
                                  "request_id": recebida.get("request_id", "")})

        if len(respostas) == 1:
            return respostas[0]
        return {"ok": True, "mensagens": respostas}

    # -------------------------------------------------------- consultores
    @api.get("/api/consultants")
    async def consultants(
        period: str = Query("30d"),
        start: str | None = None,
        end: str | None = None,
        _: str = Auth,
    ):
        since, until = range_bounds(period, config.tz, start, end)
        items = manager.consultants.list_with_stats(since, until)
        for item in items:
            item["phone_display"] = _phone_display(item.get("phone"), mask)
            if mask:
                item.pop("phone", None)
        return {"items": items, "period": {"name": period, "since": since, "until": until}}

    @api.post("/api/consultants")
    async def create_consultant(request: Request, _: str = Auth):
        data = await request.json()
        name = (data.get("name") or "").strip()
        if not name:
            raise HTTPException(status_code=400, detail="informe o nome")
        phone = "".join(ch for ch in str(data.get("phone") or "") if ch.isdigit())
        if phone and manager.consultants.exists_phone(phone):
            raise HTTPException(status_code=409, detail="já existe um consultor com esse número")
        new_id = manager.consultants.create(name, phone, notes=data.get("notes", ""))
        manager.log("INFO", "consultores", f"Consultor '{name}' cadastrado pelo painel.")
        manager.touch_metrics()
        return {"id": new_id}

    @api.put("/api/consultants/{consultant_id}")
    async def update_consultant(consultant_id: int, request: Request, _: str = Auth):
        if not manager.consultants.get(consultant_id):
            raise HTTPException(status_code=404, detail="consultor não encontrado")
        data = await request.json()
        phone = data.get("phone")
        if phone:
            digits = "".join(ch for ch in str(phone) if ch.isdigit())
            if digits and manager.consultants.exists_phone(digits, ignore_id=consultant_id):
                raise HTTPException(status_code=409, detail="já existe um consultor com esse número")
        manager.consultants.update(consultant_id, **data)
        manager.touch_metrics()
        return {"ok": True}

    @api.delete("/api/consultants/{consultant_id}")
    async def deactivate_consultant(consultant_id: int, _: str = Auth):
        if not manager.consultants.get(consultant_id):
            raise HTTPException(status_code=404, detail="consultor não encontrado")
        manager.consultants.set_active(consultant_id, False)
        manager.touch_metrics()
        return {"ok": True}

    # ---------------------------------------------------------- relatorios
    @api.get("/api/reports")
    async def reports(
        period: str = Query("30d"),
        start: str | None = None,
        end: str | None = None,
        _: str = Auth,
    ):
        return analytics.report(db, config.tz, period, start, end)

    @api.get("/api/export")
    async def export(
        background: BackgroundTasks,
        format: str = Query("csv", pattern="^(csv|xlsx|pdf)$"),
        period: str = Query("30d"),
        start: str | None = None,
        end: str | None = None,
        status: str | None = None,
        consultant: str | None = None,
        _: str = Auth,
    ):
        filters: dict[str, Any] = {"period": period, "start": start, "end": end}
        if status:
            filters["status"] = status
        if consultant:
            filters["consultant"] = consultant
        rows, _total = analytics.list_simulations(db, filters, config.tz, limit=500, offset=0)
        data = [analytics.serialize_simulation(r, mask) for r in rows]
        stamp = now_iso().replace(":", "").replace("-", "")[:15]

        if format == "csv":
            return _csv_response(data, f"simulacoes-{stamp}.csv")
        if format == "xlsx":
            return _xlsx_response(data, f"simulacoes-{stamp}.xlsx", background)
        return _pdf_response(data, f"simulacoes-{stamp}.pdf", period, background)

    # ---------------------------------------------------------------- logs
    @api.get("/api/logs")
    async def logs(
        level: str = "",
        service: str = "",
        q: str = "",
        request_id: str = "",
        consultant: str = "",
        limit: int = Query(200, ge=1, le=1000),
        offset: int = Query(0, ge=0),
        _: str = Auth,
    ):
        rows, total = analytics.list_logs(
            db, level=level, service=service, q=q, request_id=request_id,
            consultant=consultant, limit=limit, offset=offset,
        )
        return {
            "items": rows,
            "total": total,
            "services": analytics.log_services(db),
            "limit": limit,
            "offset": offset,
        }

    # ----------------------------------------------------------- tempo real
    @api.get("/api/stream")
    async def stream(request: Request, _: str = Auth):
        """Fluxo de eventos ao vivo. Rota protegida como qualquer outra."""

        async def gerar():
            fila = broadcaster.attach()
            try:
                yield _sse("hello", {"time": now_iso(), "mask_cpf": mask})
                while True:
                    try:
                        evento = await asyncio.wait_for(fila.get(), timeout=20.0)
                    except asyncio.TimeoutError:
                        # Comentario SSE: mantem a conexao viva atras de proxies
                        # sem gerar um evento para o cliente tratar.
                        yield ": ping\n\n"
                        continue
                    if await request.is_disconnected():
                        break
                    yield _sse(evento["type"], evento)
            finally:
                broadcaster.detach(fila)

        return StreamingResponse(
            gerar(),
            media_type="text/event-stream",
            headers={
                "Cache-Control": "no-cache, no-transform",
                "X-Accel-Buffering": "no",
                "Connection": "keep-alive",
            },
        )

    @api.middleware("http")
    async def no_cache_no_sniff(request: Request, call_next):
        """Painel local: o navegador precisa revalidar, e nunca adivinhar tipo.

        Sem isto, um modulo ES atualizado continuava sendo servido da cache do
        navegador depois de uma atualizacao do sistema - com o agravante de
        que so' parte dos arquivos ficava velha.
        """
        response = await call_next(request)
        path = request.url.path
        if not path.startswith("/api/"):
            response.headers["Cache-Control"] = "no-cache, must-revalidate"
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers.setdefault("Referrer-Policy", "same-origin")
        return response

    if STATIC_DIR.exists():
        api.mount("/", StaticFiles(directory=str(STATIC_DIR), html=True), name="dashboard")

    return api


# ------------------------------------------------------------------ helpers
class SseBroadcaster:
    """Uma thread lendo o barramento, N abas recebendo.

    O barramento vive em threads (WhatsApp, fila, simulador) e o servidor vive
    em asyncio. Em vez de uma thread bloqueada por aba aberta, uma unica tarefa
    drena o barramento e distribui para as filas asyncio de cada aba.
    """

    def __init__(self, hub: EventHub) -> None:
        self._hub = hub
        self._clients: set[asyncio.Queue] = set()
        self._task: asyncio.Task | None = None
        self._subscription = None

    async def start(self) -> None:
        self._subscription = self._hub.subscribe()
        self._task = asyncio.create_task(self._run())

    async def stop(self) -> None:
        if self._task:
            self._task.cancel()
            with suppress(asyncio.CancelledError):
                await self._task
            self._task = None
        if self._subscription:
            self._subscription.close()
            self._subscription = None

    def attach(self) -> asyncio.Queue:
        fila: asyncio.Queue = asyncio.Queue(maxsize=400)
        self._clients.add(fila)
        return fila

    def detach(self, fila: asyncio.Queue) -> None:
        self._clients.discard(fila)

    async def _run(self) -> None:
        loop = asyncio.get_running_loop()
        while True:
            evento = await loop.run_in_executor(None, self._subscription.get, 0.5)
            if evento is None:
                continue
            for fila in list(self._clients):
                try:
                    fila.put_nowait(evento)
                except asyncio.QueueFull:
                    # Aba lenta: perde o evento mais antigo, nunca trava o sistema.
                    with suppress(asyncio.QueueEmpty):
                        fila.get_nowait()
                    with suppress(asyncio.QueueFull):
                        fila.put_nowait(evento)


def _sse(event_type: str, data: dict) -> str:
    corpo = json.dumps(data, ensure_ascii=False, default=str)
    return f"event: {event_type}\ndata: {corpo}\n\n"


def _phone_display(phone: str | None, mask: bool) -> str:
    digits = "".join(ch for ch in (phone or "") if ch.isdigit())
    if not digits:
        return "-"
    if not mask:
        return digits
    return f"{digits[:4]}*****{digits[-2:]}" if len(digits) > 6 else "***"


def _rows_for_export(data: list[dict]) -> list[list[Any]]:
    return [[row.get(key, "") for key, _label in EXPORT_COLUMNS] for row in data]


def _csv_response(data: list[dict], filename: str) -> Response:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow([label for _key, label in EXPORT_COLUMNS])
    writer.writerows(_rows_for_export(data))
    # BOM para o Excel em pt-BR abrir com acentuacao correta.
    content = "﻿" + buffer.getvalue()
    return Response(
        content=content.encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _xlsx_response(data: list[dict], filename: str, background: BackgroundTasks) -> Response:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise HTTPException(status_code=501, detail="openpyxl não está instalado")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Simulações"
    sheet.append([label for _key, label in EXPORT_COLUMNS])
    header_fill = PatternFill("solid", fgColor="5B1A24")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="F5EFE3")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in _rows_for_export(data):
        sheet.append(row)
    for index, (_key, label) in enumerate(EXPORT_COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(
            12, min(34, len(label) + 6)
        )
    sheet.freeze_panes = "A2"

    handle = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    handle.close()
    workbook.save(handle.name)
    background.add_task(_remove_file, handle.name)
    return FileResponse(
        handle.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename,
        background=background,
    )


def _pdf_response(data: list[dict], filename: str, period: str, background: BackgroundTasks) -> Response:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        raise HTTPException(status_code=501, detail="reportlab não está instalado")

    columns = [
        ("request_id", "Solicitação", 62),
        ("created_at", "Data", 92),
        ("consultant_name", "Consultor", 100),
        ("cpf_display", "CPF", 88),
        ("bank", "Banco", 66),
        ("contract", "Contrato", 70),
        ("status_label", "Status", 62),
        ("refin", "Refin", 40),
        ("reduction_value", "Valor", 68),
        ("processing_seconds", "Tempo", 46),
    ]

    handle = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    handle.close()
    doc = SimpleDocTemplate(handle.name, pagesize=landscape(A4), title="Relatório de Simulações")
    styles = getSampleStyleSheet()

    table_data = [[label for _key, label, _w in columns]]
    for row in data[:1000]:
        table_data.append([_pdf_cell(row.get(key)) for key, _label, _w in columns])

    table = Table(table_data, repeatRows=1, colWidths=[w for _k, _l, w in columns])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5B1A24")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#F5EFE3")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C9BFAE")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F2E8")]),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    doc.build(
        [
            Paragraph("Relatório de Simulações", styles["Title"]),
            Paragraph(f"Período: {period} · Gerado em {now_iso()} · {len(data)} registro(s)",
                      styles["Normal"]),
            Spacer(1, 10),
            table,
        ]
    )
    background.add_task(_remove_file, handle.name)
    return FileResponse(handle.name, media_type="application/pdf", filename=filename,
                        background=background)


def _pdf_cell(value: Any) -> str:
    if value is None or value == "":
        return "-"
    text = str(value)
    return text if len(text) <= 26 else text[:25] + "…"


def _remove_file(path: str) -> None:
    try:
        os.unlink(path)
    except OSError:
        pass
